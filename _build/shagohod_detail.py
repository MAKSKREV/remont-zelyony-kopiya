# -*- coding: utf-8 -*-
import sys, zipfile
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход.pptx"
prs = Presentation(P)

for i in (2, 3, 5, 6, 7, 8, 10, 13, 15, 19, 22, 23, 30):
    slide = prs.slides[i - 1]
    print(f"\n===== СЛАЙД {i} (макет {slide.slide_layout.name}) =====")
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs)
                if t.strip():
                    sz = None
                    for r in p.runs:
                        if r.font.size:
                            sz = r.font.size.pt
                            break
                    print(f"   [{sz}pt] {t}")
    # заметки
    try:
        nt = slide.notes_slide.notes_text_frame.text.strip()
        if nt:
            print("   ЗАМЕТКИ:", nt[:200].replace("\n", " / "))
    except Exception:
        pass

z = zipfile.ZipFile(P)
print("\ncomments parts:", [n for n in z.namelist() if n.startswith("ppt/comments/")])
print("notesSlides parts:", len([n for n in z.namelist() if n.startswith("ppt/notesSlides/") and n.endswith(".xml")]))