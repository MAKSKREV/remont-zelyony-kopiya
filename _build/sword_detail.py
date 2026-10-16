# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\5. Инструкция для сборки Светового меча.pptx"
prs = Presentation(P)
for i in [4, 23, 24, 25, 26, 27, 38, 39, 41, 42, 43, 48]:
    slide = prs.slides[i - 1]
    print(f"\n===== SLIDE {i} =====")
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("-", sh.text_frame.text.strip().replace("\n", " // "))
        else:
            if sh.shape_type is not None and str(sh.shape_type).startswith("PICTURE"):
                pass
            try:
                if sh.shape_type == 13:
                    print("  [PIC]", sh.width, sh.height)
            except Exception:
                pass