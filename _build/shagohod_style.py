# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

print("=== LAYOUTS ===")
for i, lay in enumerate(prs.slide_layouts):
    print(f"  [{i}] {lay.name}")

def dump(idx):
    s = prs.slides[idx - 1]
    print(f"\n=== СЛАЙД {idx} (макет {s.slide_layout.name}) ===")
    for sh in s.shapes:
        print(f"  shape: {sh.shape_type} name={sh.name!r} "
              f"pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) "
              f"size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})")
        if sh.has_text_frame:
            tf = sh.text_frame
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                r = p.runs[0]
                col = None
                try:
                    col = r.font.color.rgb
                except Exception:
                    col = r.font.color.type
                print(f"      [align={p.alignment}] [font={r.font.name} {r.font.size.pt if r.font.size else None}pt "
                      f"bold={r.font.bold} color={col}] {r.text[:60]!r}")

for n in (2, 5, 34, 37):
    dump(n)
