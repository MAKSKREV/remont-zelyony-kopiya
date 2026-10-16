# -*- coding: utf-8 -*-
import sys, os, zipfile
from pptx import Presentation
from pptx.oxml.ns import qn
sys.stdout.reconfigure(encoding="utf-8")
P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\5. Инструкция для сборки Светового меча (теория КЗ).pptx"
print("size MB:", round(os.path.getsize(P)/1e6, 1))
prs = Presentation(P)
print("slides:", len(prs.slides._sldIdLst))
for i in (24, 25, 26):
    s = prs.slides[i-1]
    bu = sum(1 for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs
             if p._p.find(qn("a:pPr")) is not None and p._p.find(qn("a:pPr")).find(qn("a:buChar")) is not None)
    print(f"slide {i}: bullets={bu}, layout={s.slide_layout.name}")
# целостность остальных разделов
marks = {2:"Занятие 1", 33:"Занятие 2", 48:"Проверяем различные режимы", 51:"Готово"}
for idx, expect in marks.items():
    txt = prs.slides[idx-1].shapes[0].text_frame.text.strip() if prs.slides[idx-1].shapes else ""
    print(idx, "->", txt[:50])
z = zipfile.ZipFile(P)
print("media count:", len([n for n in z.namelist() if n.startswith("ppt/media/")]))
