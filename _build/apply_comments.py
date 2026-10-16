# -*- coding: utf-8 -*-
"""Правки по 8 комментариям школы к инструкции «Шагоход».

Комментарии (автор: Школа Изобретателей Эра Инженеров):
  слайд 5  - «меняем в основном тексте длину винтов»  -> М3х10 -> М3х16
  слайд 13 - «меняем комментарий на слайде»           -> плашка переписана под М3х16
  слайд 16 - «вставить фото или нарисовать 2 провода как этап» -> рисуем два провода
  слайд 16 - «перед этим слайдом вставить ТБ со стриппером»     -> новый слайд ТБ
  слайд 17 - «над этим слайдом должен быть слайд с ТБ паяльником» -> новый слайд ТБ
  слайд 28 - «оформляем по примеру ниже слайд»         -> перестроен как слайд «Начинаем сборку!»
  слайд 28 - «все эти комментарии подправлены. В уроках номера слайдов другие»
                                                      -> убраны номера слайдов из текста
  слайд 47 - «шрифт мелкий... обычно размер 18»        -> подписи 13pt -> 18pt
"""
import sys, os, copy

sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
SRC = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки) 3 (2).pptx")
OUT = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки 4).pptx")

FONT = "Arial"
ORANGE = RGBColor(0xFE, 0xA6, 0x21)
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x5F, 0x5F, 0x5F)
CARD = RGBColor(0xF2, 0xF3, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x00, 0x00)
WIRE_RED = RGBColor(0xD9, 0x36, 0x36)
WIRE_BLACK = RGBColor(0x2B, 0x2B, 0x2B)
COPPER = RGBColor(0xB8, 0x73, 0x33)

prs = Presentation(SRC)
LAYOUT = prs.slide_layouts[1]          # CUSTOM_2
IN = Inches


# ----------------------------------------------------------------- утилиты

def set_bullet(p, color="FEA621", indent=0.32):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    buClr.append(pPr.makeelement(qn("a:srgbClr"), {"val": color}))
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "80000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": FONT}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u25cf"}))


def para(tf, first=False, text=None, runs=None, size=18, bold=False, color=DARK,
         align=PP_ALIGN.LEFT, after=10, line=None, bullet=False):
    if runs is None:
        runs = [(text or "", {})]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    if bullet:
        set_bullet(p)
    for t, o in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", FONT)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return p


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def round_card(slide, x, y, w, h, fill=CARD, radius=0.10):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    sh.adjustments[0] = radius
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    return sh, tf


def max_shape_id(slide):
    ids = [0]
    for el in slide.shapes._spTree.iter():
        if el.tag.endswith("}cNvPr"):
            try:
                ids.append(int(el.get("id")))
            except (TypeError, ValueError):
                pass
    return max(ids)


def set_notes(slide, goal, text, note):
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    para(tf, first=True, text="Цель слайда:", size=12, bold=True, after=2)
    para(tf, text=goal, size=12, after=6)
    para(tf, text="Текст:", size=12, bold=True, after=2)
    para(tf, text=text, size=12, after=6)
    para(tf, text="Примечания:", size=12, bold=True, after=2)
    para(tf, text=note, size=12, after=0)


def replace_in_shape(shape, old, new):
    """Замена подстроки во всех абзацах, текст склеивается в первый run."""
    if not shape.has_text_frame:
        return False
    done = False
    for p in shape.text_frame.paragraphs:
        runs = p.runs
        if not runs:
            continue
        full = "".join(r.text for r in runs)
        if old in full:
            full = full.replace(old, new)
            for r in runs[1:]:
                r.text = ""
            runs[0].text = full
            done = True
    return done


# =========================================================== 1. слайды 5, 13
# комментарий «меняем в основном тексте длину винтов» + «меняем комментарий»

OLD_NOTE = ("Если винт М3х10 не дотягивается и гайка болтается - ставь М3х20 "
            "(в комплекте их 6 шт, на эти два шага нужно 4). "
            "По заявлению преподавателя нужны винты 3х16, но в комплекте их нет - "
            "уточнить у поставщика.")
NEW_NOTE = ("Основной винт здесь - М3х16. Если М3х16 нет, ставь М3х20 "
            "(в комплекте их 6 шт, на эти два шага нужно 4). "
            "В комплекте винтов М3х16 нет - докупить 4 шт: по 2 на этот слайд "
            "и на слайд «Подготавливаем детали для установки».")

for no in (5, 13):
    s = prs.slides[no - 1]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if "Винты М3х10" in txt:
            replace_in_shape(sh, "Винты М3х10", "Винты М3х16")
            print(f"[{no}] винты -> М3х16")
        if "Не дотягивается" in txt:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("М3х10", "М3х16")
                    r.font.size = Pt(14)
            print(f"[{no}] плашка -> М3х16, 14pt")
    ntf = s.notes_slide.notes_text_frame
    if OLD_NOTE in ntf.text:
        ntf.text = ntf.text.replace(OLD_NOTE, NEW_NOTE)
        print(f"[{no}] заметки обновлены")

# =========================================================== 2. слайд 16
# комментарий «вставить фото или нарисовать 2 провода как этап»

s16 = prs.slides[15]

line = s16.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, IN(0.9), IN(1.55), IN(8.65), IN(1.55))
line.line.color.rgb = GRAY
line.line.width = Pt(1.5)
for x in (0.9, 8.65):
    tick = s16.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, IN(x), IN(1.45), IN(x), IN(1.65))
    tick.line.color.rgb = GRAY
    tick.line.width = Pt(1.5)

pill, ptf = round_card(s16, 3.5, 1.30, 2.5, 0.5, fill=WHITE, radius=0.5)
para(ptf, first=True, text="12-16 см", size=18, bold=True, color=ORANGE,
     align=PP_ALIGN.CENTER, after=0)

wires = [
    (2.10, WIRE_RED, "1"),
    (3.10, WIRE_BLACK, "2"),
]
for y, color, num in wires:
    body = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(0.9), IN(y), IN(7.1), IN(0.26))
    body.adjustments[0] = 0.5
    body.shadow.inherit = False
    body.fill.solid()
    body.fill.fore_color.rgb = color
    body.line.fill.background()
    bare = s16.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(8.0), IN(y + 0.07), IN(0.65), IN(0.12))
    bare.shadow.inherit = False
    bare.fill.solid()
    bare.fill.fore_color.rgb = COPPER
    bare.line.fill.background()
    dot = s16.shapes.add_shape(MSO_SHAPE.OVAL, IN(0.40), IN(y - 0.04), IN(0.34), IN(0.34))
    dot.shadow.inherit = False
    dot.fill.solid()
    dot.fill.fore_color.rgb = ORANGE
    dot.line.fill.background()
    tfd = dot.text_frame
    tfd.margin_left = tfd.margin_right = 0
    tfd.margin_top = tfd.margin_bottom = 0
    tfd.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tfd, first=True, text=num, size=14, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, after=0)

_, ctf = round_card(s16, 0.5, 3.85, 9.0, 0.95, fill=CARD, radius=0.10)
ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(ctf, first=True, text="Концы зачисти на 5 мм с каждой стороны: жилы целые и блестят.",
     size=16, align=PP_ALIGN.CENTER, after=4)
para(ctf, text="Оба провода одинаковой длины - оба пойдут на двигатель.",
     size=14, color=GRAY, align=PP_ALIGN.CENTER, after=0)

set_notes(s16,
          "Показать, как подготовить два провода перед пайкой.",
          "Нарезаем два провода по 12-16 см, зачищаем концы на 5 мм и проверяем жилы.",
          "Перед этим слайдом идёт техника безопасности со стриппером (съемником изоляции), "
          "после - техника безопасности с паяльником. Длина проводов указана прямо на слайде, "
          "без ссылок на другие номера слайдов.")
print("[16] нарисованы два провода + размер")

# =========================================================== 3. слайд 17
# мелкая опечатка в подписи + заметки под новые слайды

s17 = prs.slides[16]
for sh in s17.shapes:
    if replace_in_shape(sh, "Эти провод длиной", "Эти провода длиной"):
        print("[17] подпись исправлена: «Эти провода длиной 12-16 см»")
set_notes(s17,
          "Показать начало пайки: два провода, кнопка и плата.",
          "Паяем схему, как показано на фото.",
          "Перед этим слайдом идёт техника безопасности с паяльником. "
          "Провода подготовлены на предыдущем слайде, длина 12-16 см. "
          "Ссылки на номера слайдов не давать - в уроках нумерация другая.")

# =========================================================== 4. слайд 26
# подпись 13pt -> 14pt, заметки без номеров слайдов

s26 = prs.slides[25]
for sh in s26.shapes:
    if sh.has_text_frame and "Кнопку и электронику" in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(14)
        sh.height = Inches(0.8)
        print("[26] подпись 13 -> 14pt")
set_notes(s26,
          "Показать установку электроники в корпус.",
          "Электроника ставится так же, как кнопка и схема на слайде «Начнем пайку».",
          "Раньше здесь была ссылка без номера - ребенок не понимал, куда смотреть. "
          "Теперь ссылка дана названием слайда: номера слайдов в уроках отличаются, "
          "поэтому в тексте номера не указываем. Рядом добавлено фото.")

# =========================================================== 5. слайд 28
# «оформляем по примеру ниже слайд» + убрать номера слайдов из текста

s28 = prs.slides[27]
s29 = prs.slides[28]

card28 = None
hint28 = None
for sh in s28.shapes:
    if sh.shape_id == 424:
        card28 = sh
    elif sh.shape_id == 425:
        hint28 = sh
if card28 is not None:
    card28._element.getparent().remove(card28._element)
    print("[28] серая карточка удалена")

# --- список: копия плейсхолдера со слайда 29 (точно в их стиле)
src_list = None
src_pic = None
for sh in s29.shapes:
    if sh.shape_id == 431:
        src_list = sh
    elif sh.shape_id == 432:
        src_pic = sh

new_id = max_shape_id(s28) + 1

list_el = copy.deepcopy(src_list._element)
cnv = list_el.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
cnv.set("id", str(new_id))
cnv.set("name", f"Google Shape;{new_id};copy28")
spPr = list_el.find(qn("p:spPr"))
xfrm = spPr.find(qn("a:xfrm"))
xfrm.find(qn("a:off")).set("x", str(Inches(0.35)))
xfrm.find(qn("a:off")).set("y", str(Inches(0.72)))
xfrm.find(qn("a:ext")).set("cx", str(Inches(4.50)))
xfrm.find(qn("a:ext")).set("cy", str(Inches(3.95)))

txBody = list_el.find(qn("p:txBody"))
ps = txBody.findall(qn("a:p"))
tmpl_bullet = copy.deepcopy(ps[1])


def set_p_text(p_el, text):
    runs = p_el.findall(qn("a:r"))
    for extra in runs[1:]:
        p_el.remove(extra)
    if runs:
        t = runs[0].find(qn("a:t"))
        t.text = text


set_p_text(ps[0], "Что понадобится на занятии 2:")
for p_el in ps[1:]:
    txBody.remove(p_el)
ITEMS28 = [
    "Винты М3х20 - 2 шт. и гайки М3 - 2 шт. - детали с винтами",
    "Гайки М3 - 8 шт.: стенка - 2, корпус - 3, вторая деталь - 3",
    "Деревянные шайбы - на 3 места: вал двигателя, крепление, вторая деталь",
    "Электроника с занятия 1: батарейки, кнопка, двигатель, провода",
    "Клей и инструмент - как на занятии 1",
]
for it in ITEMS28:
    newp = copy.deepcopy(tmpl_bullet)
    set_p_text(newp, it)
    txBody.append(newp)
s28.shapes._spTree.append(list_el)
print("[28] список перестроен по образцу слайда «Начинаем сборку!»")

# --- фото: копия со слайда 29 (с сохранением кадрирования)
pic_el = copy.deepcopy(src_pic._element)
cnv = pic_el.find(qn("p:nvPicPr")).find(qn("p:cNvPr"))
cnv.set("id", str(new_id + 1))
cnv.set("name", f"Google Shape;{new_id + 1};copy28")
old_rid = pic_el.find(qn("p:blipFill")).find(qn("a:blip")).get(qn("r:embed"))
img_part = s29.part.related_part(old_rid)
new_rid = s28.part.relate_to(img_part, RT.IMAGE)
pic_el.find(qn("p:blipFill")).find(qn("a:blip")).set(qn("r:embed"), new_rid)
pxfrm = pic_el.find(qn("p:spPr")).find(qn("a:xfrm"))
pxfrm.find(qn("a:off")).set("x", str(Inches(5.00)))
pxfrm.find(qn("a:off")).set("y", str(Inches(0.99)))
s28.shapes._spTree.append(pic_el)
print("[28] фото скопировано со слайда «Начинаем сборку!»")

# --- подвал-подсказка без номеров слайдов
hint28.left = IN(0.35)
hint28.top = IN(4.85)
hint28.width = IN(9.30)
hint28.height = IN(0.65)
htf = hint28.text_frame
htf.word_wrap = True
ps = htf.paragraphs
set_runs = [
    ("Разложи детали по группам перед занятием - так соберешь быстрее.", 15, ORANGE, True),
    ("Шагоход не поедет? Разбор неисправностей - в конце инструкции.", 14, GRAY, False),
]
for p, (txt, sz, col, bd) in zip(ps, set_runs):
    for r in p.runs[1:]:
        r.text = ""
    if p.runs:
        p.runs[0].text = txt
        p.runs[0].font.size = Pt(sz)
        p.runs[0].font.color.rgb = col
        p.runs[0].font.bold = bd
print("[28] подсказка переписана без номеров слайдов")

set_notes(s28,
          "Дать ребенку заранее подготовить все детали на занятие 2.",
          "Список деталей перед началом занятия.",
          "Оформление - по образцу следующего слайда («Начинаем сборку!»). "
          "Итого гаек М3 нужно 10 шт.: стенка - 2, детали с винтами - 2, корпус - 3, "
          "вторая деталь - 3. На слайде «Начинаем сборку!» указано 8 гаек - там не посчитаны "
          "2 гайки для деталей с винтами. Номера слайдов в тексте не указываем: "
          "в уроках нумерация другая.")

# =========================================================== 6. слайд 47
# «шрифт мелкий... обычно размер 18»

s47 = prs.slides[46]
CAPS = {
    570: (0.40, 2.32), 572: (5.15, 2.32),
    574: (0.40, 4.56), 576: (5.15, 4.56),
}
for sh in s47.shapes:
    if sh.shape_id in CAPS and sh.has_text_frame:
        x, y = CAPS[sh.shape_id]
        sh.left, sh.top = IN(x), IN(y)
        sh.width, sh.height = IN(4.60), IN(0.55)
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18)
        print(f"[47] подпись {sh.shape_id}: 13 -> 18pt")
set_notes(s47,
          "Показать готовые узлы, которые проверяем перед первым запуском.",
          "Четыре фото из инструкции с короткими подписями.",
          "Фото взяты из других слайдов этой же инструкции "
          "(электроника, приклеивание, стенка, шайбы). Можно использовать как памятку "
          "- распечатать отдельно. Номера слайдов не указываем: в уроках нумерация другая.")

# =========================================================== 7. новые слайды ТБ

def tb_slide(title, banner, bullets, danger, icon_caption, goal, text, note):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    s.shapes.title.text_frame.text = title

    sh, tf = round_card(s, 0.40, 0.66, 9.20, 0.58, fill=ORANGE, radius=0.14)
    para(tf, first=True, text=banner, size=17, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, after=0)

    tfb = textbox(s, 0.50, 1.45, 6.20, 2.95)
    for i, b in enumerate(bullets):
        para(tfb, first=(i == 0), text=b, size=17, after=9, line=1.05, bullet=True)

    tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, IN(7.10), IN(1.55), IN(2.10), IN(1.85))
    tri.shadow.inherit = False
    tri.fill.solid()
    tri.fill.fore_color.rgb = ORANGE
    tri.line.fill.background()
    ttf = tri.text_frame
    ttf.word_wrap = False
    ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ttf, first=True, text="!", size=44, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, after=0)
    tfc = textbox(s, 6.75, 3.50, 2.80, 0.45)
    para(tfc, first=True, text=icon_caption, size=14, color=GRAY,
         align=PP_ALIGN.CENTER, after=0)

    sh2, tf2 = round_card(s, 0.40, 4.55, 9.20, 0.60, fill=RED, radius=0.14)
    para(tf2, first=True, text=danger, size=17, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, after=0)

    set_notes(s, goal, text, note)
    return s


s_tb_stripper = tb_slide(
    "Техника безопасности: стриппер",
    "Стриппер снимает только изоляцию - держи провод пальцами позади лезвия.",
    [
        "Работаем за столом - не на коленях и не на ходу.",
        "Провод держим пальцами позади лезвия, а не перед ним.",
        "Направляй стриппер от себя и снимай только изоляцию - жилу не режь.",
        "Проверь конец после зачистки: жилы целые и блестят.",
        "После работы убери стриппер на место, лезвие закрыто.",
    ],
    "Без взрослого стриппер не берем - и никогда не режем на весу.",
    "Осторожно: лезвие",
    "Провести правила безопасности перед подготовкой проводов.",
    "Пять простых правил работы со стриппером (съемником изоляции).",
    "Стриппер в этом проекте заменяет нож: снимает изоляцию с проводов 12-16 см. "
    "Следующий слайд - подготовка двух проводов. Правила короткие - "
    "назвать их вслух и показать движение руки.")

s_tb_iron = tb_slide(
    "Техника безопасности: паяльник",
    "Жало паяльника нагревается до 300-400 градусов - трогать нельзя!",
    [
        "Паяльник ставим только в держатель - не на стол и не на бумагу.",
        "Подноси провод к жалу, а не жало к руке; волосы и рукава убраны.",
        "Дым от припоя не вдыхай - работай в проветриваемой комнате.",
        "Отпаял - выключи из розетки и убери в держатель.",
        "Жало горячее еще 10 минут после выключения - не трогай.",
    ],
    "Обжегся - подержи руку под холодной водой и скажи взрослому.",
    "Ожог! 300-400 градусов",
    "Провести правила безопасности перед пайкой.",
    "Пять правил работы с паяльником и что делать при ожоге.",
    "Паяльник применяется на слайде «Начнем пайку». Показать держатель, "
    "проверить, что волосы и рукава не свисают, проветрить комнату. "
    "При ожоге - холодная вода и взрослый.")

print("[+] слайд ТБ со стриппером и слайд ТБ с паяльником созданы")

# =========================================================== 7b. слайд с гифкой
# заглушка «Добавить гифку работы» при уже вставленной гифке

s_gif = prs.slides[43]
for sh in s_gif.shapes:
    if replace_in_shape(sh, "Добавить гифку работы", "Шагоход в движении"):
        print("[44] заглушка -> «Шагоход в движении»")
set_notes(s_gif,
          "Показать готовый результат - как шагоход ходит.",
          "Гифка работы шагохода.",
          "Гифка (1,7 МБ) уже вставлена в слайд - подпись «Добавить гифку работы» "
          "была незакрытой задачей, заменена на название слайда.")

# =========================================================== 8. перестановка

def move_last_to(prs, index):
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    moved, rest = ids[-1], ids[:-1]
    new_ids = rest[:index] + [moved] + rest[index:]
    for el in list(sldIdLst):
        sldIdLst.remove(el)
    for el in new_ids:
        sldIdLst.append(el)


# в конце списка сейчас: [... 48 старых ..., ТБ-стриппер, ТБ-паяльник]
# хотим: 16 = ТБ со стриппером, 17 = «Готовим два провода», 18 = ТБ с паяльником,
#        19 = «Начнем пайку»
move_last_to(prs, 16)   # паяльник - перед «Начнем пайку» (индекс 16)
move_last_to(prs, 15)   # стриппер - перед «Готовим два провода» (индекс 15)

prs.save(OUT)
print("\nсохранено:", OUT)
print("слайдов:", len(prs.slides._sldIdLst))
