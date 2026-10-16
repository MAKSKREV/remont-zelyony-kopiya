# -*- coding: utf-8 -*-
"""Легенда (слайды 2-3), ссылки «слайд» по всему деку, шрифты <14."""
import sys, os, re
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu, Pt

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
P = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки) 3 (2).pptx")
prs = Presentation(P)

print("=== слайды 2, 3 (легенда) ===")
for no in (2, 3):
    s = prs.slides[no - 1]
    print(f"--- {no} ---")
    for sh in s.shapes:
        if sh.shape_type == 6:
            print(f"  ГРУППА '{sh.name}' ({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f})")
            for sub in sh.shapes:
                t = sub.text_frame.text if sub.has_text_frame else ""
                print(f"    {sub.shape_type} '{sub.name}' fill=? text={t!r}")
                try:
                    from lxml import etree
                    print("      ", etree.tostring(sub._element, pretty_print=True).decode()[:900].replace("\n", " "))
                except Exception as e:
                    print("      err", e)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            print(f"  TEXT: {sh.text_frame.text.strip()!r}")

print("\n=== все тексты со словом «слайд» (видимый текст) ===")
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            for line in t.split("\n"):
                if "слайд" in line.lower():
                    print(f"  [{i}] {line.strip()!r}")

print("\n=== заметки со словом «слайд» ===")
for i, s in enumerate(prs.slides, 1):
    try:
        t = s.notes_slide.notes_text_frame.text
    except Exception:
        continue
    for line in t.split("\n"):
        if "слайд" in line.lower():
            print(f"  [заметки {i}] {line.strip()!r}")

print("\n=== видимые шрифты < 14pt ===")
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and r.font.size.pt < 14 and r.text.strip():
                    print(f"  [{i}] {r.font.size.pt}pt: {r.text.strip()[:80]!r}")

print("\n=== слайд 48 ===")
s = prs.slides[47]
print("layout:", s.slide_layout.name, "shapes:", len(s.shapes))
for sh in s.shapes:
    print("  ", sh.shape_type, sh.name, sh.has_text_frame and sh.text_frame.text)
