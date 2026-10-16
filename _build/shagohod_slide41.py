# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
s = prs.slides[40]  # слайд 41
for sh in s.shapes:
    if sh.shape_type == 13:
        img = sh.image
        print("файл:", img.filename, "| ContentType:", img.content_type)
        print("размер в байтах:", len(img.blob), "| px:", img.size)
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("текст фигуры:", sh.text_frame.text.strip(), "| left", sh.left, "| is_placeholder", sh.is_placeholder)