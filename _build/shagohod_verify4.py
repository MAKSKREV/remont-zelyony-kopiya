# -*- coding: utf-8 -*-
import sys, re, os, zipfile
from collections import defaultdict
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
print("слайдов:", len(prs.slides._sldIdLst), "| MB:", round(os.path.getsize(P)/1e6, 2))
z = zipfile.ZipFile(P)
print("zip:", "ок" if z.testzip() is None else "ОШИБКА")

print("\n=== структура ===")
for i, s in enumerate(prs.slides, 1):
    t = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
    if not t:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = "[" + sh.text_frame.text.strip().split("\n")[0][:40] + "]"
                break
    print(f"[{i:2}] {t[:60]}")

print("\n=== сверка комплекта (без служебных слайдов 4 и 27) ===")
use = defaultdict(int)
for i, s in enumerate(prs.slides, 1):
    if i in (4, 27):
        continue
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            m = re.match(r"^(Винты|Гайки)\s+(М3\S*)\s*-\s*(\d+)\s*шт", t, re.I)
            if m:
                use[(m.group(1), m.group(2))] += int(m.group(3))
for k in sorted(use):
    print("  ", k[0], k[1], use[k], "шт")
print("  всего гаек М3:", use[("Гайки", "М3")], "из 23 в комплекте ->",
      "хватает" if use[("Гайки", "М3")] <= 23 else "НЕ ХВАТАЕТ")

print("\n=== все ссылки на слайды ===")
REF = re.compile(r"слайд[а-я]*\s+\d+(?:\s*[-–]\s*\d+)?(?:\s*(?:,|и)\s*\d+)*", re.I)
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for m in REF.finditer(sh.text_frame.text):
                print(f"  слайд {i:2}: {m.group(0)}")
    try:
        for m in REF.finditer(s.notes_slide.notes_text_frame.text):
            print(f"  заметки {i:2}: {m.group(0)}")
    except Exception:
        pass