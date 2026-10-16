# -*- coding: utf-8 -*-
import sys, shutil, os
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Пример.pptx"
SCR = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\_build\test.pptx"

def delete_all_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn('r:id'))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)

print("src size:", os.path.getsize(SRC))
shutil.copy2(SRC, SCR)
prs = Presentation(SCR)
print("before sldIdLst:", len(prs.slides._sldIdLst), "| layouts:", len(prs.slide_masters[0].slide_layouts))
delete_all_slides(prs)
prs.save(SCR)
print("saved size:", os.path.getsize(SCR))
prs2 = Presentation(SCR)
print("reopened sldIdLst:", len(prs2.slides._sldIdLst), "| layouts:", len(prs2.slide_masters[0].slide_layouts))