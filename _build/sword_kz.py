# -*- coding: utf-8 -*-
"""Вставка блоков про КЗ и поиск ошибок в «Инструкцию для сборки Светового меча».
Новые слайды размещаются сразу после слайда 23 «Электрическая схема»."""
import sys, shutil

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

SRC = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\5. Инструкция для сборки Светового меча.pptx"
OUT = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\5. Инструкция для сборки Светового меча (теория КЗ).pptx"

ORANGE = RGBColor(0xFE, 0xA6, 0x21)
DARK = RGBColor(0x26, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x00, 0x00)
CREAM = RGBColor(0xFF, 0xF4, 0xE0)
FONT = "Arial"
W, H = 9144000, 5143500


def set_bullet(p, color="FEA621", char="●", sz=1800):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', '457200')
    pPr.set('indent', '-342900')
    for tag in ('a:buClrTx', 'a:buClr', 'a:buSzTx', 'a:buSzPct', 'a:buSzPts',
                'a:buFontTx', 'a:buFont', 'a:buNone', 'a:buAutoNum', 'a:buChar'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pos = 0
    for child in pPr:
        if child.tag in (qn('a:lnSpc'), qn('a:spcBef'), qn('a:spcAft')):
            pos += 1
        else:
            break
    def mk(tag, attrs=None, inner=()):
        el = pPr.makeelement(qn(tag), attrs or {})
        for t2, a2 in inner:
            el.append(mk(t2, a2))
        return el
    items = [
        mk('a:buClr', inner=(('a:srgbClr', {'val': color}),)),
        mk('a:buSzPts', {'val': str(sz)}),
        mk('a:buFont', {'typeface': 'Arial'}),
        mk('a:buChar', {'char': char}),
    ]
    for el in items:
        pPr.insert(pos, el)
        pos += 1


def add_title(slide, text, size=23):
    tb = slide.shapes.add_textbox(0, 0, W, 514200)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Emu(91425)
    tf.margin_top = Emu(91425)
    tf.margin_right = Emu(91425)
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = False
    f.color.rgb = DARK
    return tf


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def para(tf, first=False, text=None, runs=None, size=18, bold=False, color=DARK,
         align=PP_ALIGN.LEFT, after=10, line=None, bullet=False, bullet_color="FEA621"):
    if runs is None:
        runs = [(text or "", {})]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    if after is not None:
        p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    if bullet:
        set_bullet(p, color=bullet_color)
    for t, o in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get('font', FONT)
        f.size = Pt(o.get('size', size))
        f.bold = o.get('bold', bold)
        f.italic = o.get('italic', False)
        f.color.rgb = o.get('color', color)
    return p


def card(slide, x, y, w, h, fill=WHITE, line=None, radius=0.12, line_w=1.5):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = radius
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Pt(18)
    tf.margin_top = tf.margin_bottom = Pt(6)
    return sp, tf


def move_slides_after(prs, anchor_slide_no, count):
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    moved = ids[-count:]
    rest = ids[:-count]
    insert_at = anchor_slide_no  # insert before element index anchor_slide_no (0-based)
    new_ids = rest[:insert_at] + moved + rest[insert_at:]
    for el in list(sldIdLst):
        sldIdLst.remove(el)
    for el in new_ids:
        sldIdLst.append(el)


def new_content_slide(prs):
    layout = prs.slide_masters[0].slide_layouts[1]  # CUSTOM_2
    return prs.slides.add_slide(layout)


def main():
    shutil.copy2(SRC, OUT)
    prs = Presentation(OUT)
    n_before = len(prs.slides._sldIdLst)

    # ---------- Слайд 1: Короткое замыкание ----------
    s = new_content_slide(prs)
    add_title(s, "Короткое замыкание (КЗ)")
    sp, tf = card(s, 400000, 660000, 8300000, 400000, fill=ORANGE, line=None, radius=0.12)
    para(tf, first=True,
         runs=[("КЗ — короткое замыкание: ", {'bold': True, 'size': 20, 'color': WHITE}),
               ("плюс (+) и минус (−) соединились напрямую, в обход светодиодной ленты.",
                {'size': 18, 'color': WHITE})], align=PP_ALIGN.CENTER, after=0)
    tf = textbox(s, 500000, 1220000, 7600000, 1900000)
    para(tf, first=True, runs=[("Ток идёт не через ленту, ", {'bold': True, 'size': 18}),
                               ("а напрямую между плюсом и минусом.", {'size': 18})],
         bullet=True, after=14, line=1.1)
    para(tf, runs=[("Лента гаснет или светит тускло, ", {'bold': True, 'size': 18}),
                   ("батарейка греется и быстро садится.", {'size': 18})],
         bullet=True, after=14, line=1.1)
    para(tf, runs=[("Пайка и провода в месте замыкания ", {'bold': True, 'size': 18}),
                   ("могут оплавиться.", {'size': 18})], bullet=True, after=14, line=1.1)
    sp, tf = card(s, 400000, 3220000, 8300000, 560000, fill=CREAM, line=ORANGE, radius=0.1)
    para(tf, first=True, runs=[("Когда возникает в нашем мече: ", {'bold': True, 'size': 18}),
                               ("оголённые провода касаются друг друга, застряли в клемме "
                                "или плохо изолированы.", {'size': 18})],
         align=PP_ALIGN.CENTER, after=0)
    sp, tf = card(s, 400000, 3980000, 8300000, 380000, fill=RED, line=None, radius=0.12)
    para(tf, first=True, text="Главное правило: плюс и минус никогда не должны касаться друг друга!",
         size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, after=0)

    # ---------- Слайд 2: Как искать ошибки ----------
    s = new_content_slide(prs)
    add_title(s, "Как искать ошибки в схеме")
    steps = [
        ("1. Начни с батарейки. ", "Проверь заряд и контакты в клеммнике."),
        ("2. Проверь изоляцию. ", "Огляди все спайки и клеммы: не касаются ли плюс и минус."),
        ("3. Иди по схеме по порядку. ", "Батарейка → клеммник → коннектор → лента. Проверь каждый узел глазами."),
        ("4. Прозвони цепь. ", "С преподавателем прозвони мультиметром: где нет сигнала — там обрыв."),
        ("5. Перепаяй аккуратно. ", "Часто ошибка уходит после перепайки одного узла."),
    ]
    tf = textbox(s, 500000, 680000, 7600000, 3300000)
    for i, (lead, body) in enumerate(steps):
        para(tf, first=(i == 0), runs=[(lead, {'bold': True, 'size': 19}),
                                       (body, {'size': 18})],
             bullet=True, bullet_color="FEA621", after=18, line=1.1)
    sp, tf = card(s, 400000, 4000000, 8300000, 460000, fill=CREAM, line=ORANGE, radius=0.1)
    para(tf, first=True, text="Правило: исправляем по одному узлу и сразу проверяем результат — "
                              "не ждём, когда засветится всё сразу.",
         size=18, bold=True, color=DARK, align=PP_ALIGN.CENTER, after=0)

    # ---------- Слайд 3: Разбор «цвета не переключаются» ----------
    s = new_content_slide(prs)
    add_title(s, "Разбор: не переключаются все цвета")
    sp, tf = card(s, 400000, 660000, 8300000, 380000, fill=ORANGE, line=None, radius=0.12)
    para(tf, first=True,
         runs=[("Ситуация: ", {'bold': True, 'size': 20, 'color': WHITE}),
               ("лента горит только одним цветом, а кнопки не включают остальные.",
                {'size': 18, 'color': WHITE})], align=PP_ALIGN.CENTER, after=0)
    tf = textbox(s, 500000, 1160000, 7600000, 2600000)
    para(tf, first=True, runs=[("Три провода ленты отвечают за три цвета: ", {'size': 18}),
                               ("R (красный), G (зелёный), B (синий) + общий минус.",
                                {'bold': True, 'size': 18})], bullet=True, after=14, line=1.1)
    para(tf, runs=[("Не горит один цвет → ", {'bold': True, 'size': 18}),
                   ("обрыв или КЗ на его проводе. Проверь именно R, G или B.",
                    {'size': 18})], bullet=True, after=14, line=1.1)
    para(tf, runs=[("Не горит ничего → ", {'bold': True, 'size': 18}),
                   ("проверь общий минус (−) и питание.", {'size': 18})],
         bullet=True, after=14, line=1.1)
    para(tf, runs=[("Кнопки — просто выключатели. ", {'bold': True, 'size': 18}),
                   ("Если кнопка не пропаяна, цвет до неё «не доходит».", {'size': 18})],
         bullet=True, after=14, line=1.1)
    sp, tf = card(s, 400000, 3860000, 8300000, 700000, fill=CREAM, line=ORANGE, radius=0.1)
    para(tf, first=True, runs=[("Примени алгоритм «Как искать ошибки»: ", {'bold': True, 'size': 18}),
                               ("иди по цепочке по одному узлу.", {'size': 18})],
         align=PP_ALIGN.CENTER, after=6)
    para(tf, runs=[("Проверь: ", {'bold': True, 'size': 18}),
                   ("3 провода коннектора → пайки → переключатели → общий минус.",
                    {'size': 17})], align=PP_ALIGN.CENTER, after=0)

    move_slides_after(prs, 23, 3)
    prs.save(OUT)
    print("saved:", OUT)
    print("slides:", n_before, "->", len(prs.slides._sldIdLst))

    prs2 = Presentation(OUT)
    for i, slide in enumerate(prs2.slides, 1):
        if i in (22, 23, 24, 25, 26, 27):
            txts = [sh.text_frame.text.strip().replace("\n", " / ")[:60]
                    for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            print(i, "::", " | ".join(txts[:2]))


if __name__ == "__main__":
    main()