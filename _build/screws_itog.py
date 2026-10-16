# -*- coding: utf-8 -*-
"""«шагоход итог»: сколько и каких винтов/гаек используется."""
import sys, os, re
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
P = os.path.join(BASE, "шагоход итог.pptx")
prs = Presentation(P)
print("слайдов:", len(prs.slides))

# структура
print("\n=== структура ===")
for i, s in enumerate(prs.slides, 1):
    t = ""
    for sh in s.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 0 and sh.has_text_frame:
            t = sh.text_frame.text.strip().replace("\n", " "); break
    if not t:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().replace("\n", " ")[:60]; break
    print(f"{i:3}. {t}")

# таблица: слайд -> строки с винтами/гайками
print("\n=== построчно: винты/гайки в видимом тексте ===")
tot = {}
nut_tot = 0
rows = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for line in sh.text_frame.text.split("\n"):
            line = line.strip()
            m = re.search(r"Винты?\s+М3\s*х\s*(\d+)\s*[-–]\s*(\d+)", line, re.I)
            if m:
                size, cnt = m.group(1), int(m.group(2))
                tot[size] = tot.get(size, 0) + cnt
                rows.append((i, size, cnt, line))
            g = re.search(r"Гайки?\s+М3\s*[-–]\s*(\d+)", line, re.I)
            if g:
                nut_tot += int(g.group(1))
                rows.append((i, "гайка", int(g.group(1)), line))
for r in rows:
    print(f"  [{r[0]:>2}] М3х{r[1]}: {r[2]}  <- {r[3]}")

print("\n=== ИТОГО винтов по слайдам ===")
for k in sorted(tot, key=int):
    print(f"  М3х{k}: {tot[k]} шт")
print(f"  ВСЕГО винтов: {sum(tot.values())} шт")
print(f"  Гаек М3 (по подписям на слайдах): {nut_tot} шт")

# любые другие упоминания винтов
print("\n=== прочие упоминания винт/гайк ===")
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for line in sh.text_frame.text.split("\n"):
                if re.search(r"винт|гайк", line, re.I) and not re.search(r"М3\s*х\s*\d+\s*[-–]|Гайки?\s+М3\s*[-–]\s*\d", line):
                    print(f"  [{i}] {line.strip()[:110]}")
