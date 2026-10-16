# -*- coding: utf-8 -*-
"""Replace PowerPoint comments with speaker notes (заметки докладчика)."""
import os, zipfile
from lxml import etree
from add_comments import COMMENTS
from pptx import Presentation

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
SRC  = os.path.join(BASE, "TRIZ-master-klass для преподавателя.pptx")
CLEAN = os.path.join(BASE, "_build", "clean.pptx")
OUT1 = os.path.join(BASE, "TRIZ-master-klass.pptx")
OUT2 = os.path.join(BASE, "TRIZ-master-klass для преподавателя.pptx")

NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_COMMENTS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
REL_AUTHORS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors"


def strip_comments(src, dst):
    with zipfile.ZipFile(src) as zin:
        data = {n: zin.read(n) for n in zin.namelist()}

    ct = etree.fromstring(data["[Content_Types].xml"])
    for o in list(ct):
        if o.tag == f"{{{NS_CT}}}Override":
            pn = o.get("PartName", "")
            if pn.startswith("/ppt/comments/comment") or pn == "/ppt/commentAuthors.xml":
                ct.remove(o)
    data["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True,
                                                 encoding="UTF-8", standalone=True)

    rp = "ppt/_rels/presentation.xml.rels"
    rels = etree.fromstring(data[rp])
    for r in list(rels):
        if r.get("Type") == REL_AUTHORS:
            rels.remove(r)
    data[rp] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    for name in [n for n in data if n.startswith("ppt/slides/_rels/slide") and n.endswith(".rels")]:
        rels = etree.fromstring(data[name])
        for r in list(rels):
            if r.get("Type") == REL_COMMENTS:
                rels.remove(r)
        data[name] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    for n in list(data):
        if n.startswith("ppt/comments/comment") or n == "ppt/commentAuthors.xml":
            del data[n]

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for n, b in data.items():
            zo.writestr(n, b)
    print("comments stripped ->", dst)


def set_notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.text = ""
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line


strip_comments(SRC, CLEAN)

prs = Presentation(CLEAN)
for k, slide in enumerate(prs.slides, 1):
    set_notes(slide, COMMENTS[k])
prs.save(OUT1)
prs.save(OUT2)
print("speaker notes added to", len(prs.slides._sldIdLst), "slides")

for _out in (OUT1, OUT2):
    p2 = Presentation(_out)
    n1 = p2.slides[0].notes_slide.notes_text_frame.text
    n9 = p2.slides[8].notes_slide.notes_text_frame.text
    print(os.path.basename(_out), "| slide1 notes:", repr(n1[:50]))

import zipfile as zz
with zz.ZipFile(OUT1) as z:
    print("comments left:", [n for n in z.namelist() if "comments" in n or "commentAuthors" in n])