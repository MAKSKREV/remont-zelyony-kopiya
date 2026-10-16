# -*- coding: utf-8 -*-
"""Правка текста: слайд 7 «еще четыре таких детали» -> «еще три таких детали»."""
import sys, shutil
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

SRC = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход.pptx"
OUT = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"

shutil.copy2(SRC, OUT)
prs = Presentation(OUT)

slide = prs.slides[6]  # слайд 7
changed = []
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        runs = p.runs
        if not runs:
            continue
        full = "".join(r.text for r in runs)
        if "четыре" in full:
            # заменяем только в ране, где встречается слово
            for r in runs:
                if "четыре" in r.text:
                    old = r.text
                    r.text = old.replace("четыре", "три")
                    changed.append((old, r.text))
prs.save(OUT)
print("изменено фрагментов:", len(changed))
for a, b in changed:
    print("  было :", a)
    print("  стало:", b)

# проверка
prs2 = Presentation(OUT)
for i in (6, 7, 8):
    s = prs2.slides[i - 1]
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().replace("\n", " | ")
            if "Собираем" in t or "Винты" in t:
                print(f"слайд {i}:", t)