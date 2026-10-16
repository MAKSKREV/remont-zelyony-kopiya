# -*- coding: utf-8 -*-
import sys, re
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from lxml import etree

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

lay = prs.slide_layouts[1]
print("=== LAYOUT CUSTOM_2 : title ph XML ===")
for ph in lay.placeholders:
    if ph.placeholder_format.idx == 0:
        print(etree.tostring(ph._element, pretty_print=True).decode()[:2500])

print("=== LAYOUT CUSTOM_2 : body ph XML (idx=2) ===")
for ph in lay.placeholders:
    if ph.placeholder_format.idx == 2:
        print(etree.tostring(ph._element, pretty_print=True).decode()[:2000])

print("=== LAYOUT CUSTOM_2 : прочие фигуры ===")
for sh in lay.shapes:
    if not sh.is_placeholder:
        print("  ", sh.shape_type, repr(sh.name))
        xml = etree.tostring(sh._element).decode()
        fills = re.findall(r'(?:srgbClr val="([0-9A-Fa-f]{6})"|schemeClr val="(\w+)")', xml)
        print("      цвета:", fills[:10])

theme = prs.slide_masters[0].part.part_related_by(
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
tx = theme.blob.decode("utf-8", "ignore")
m = re.search(r"<a:clrScheme.*?</a:clrScheme>", tx, re.S)
print("=== ТЕМА ЦВЕТА ===")
print(m.group(0) if m else "не найдено")
m2 = re.search(r"<a:fontScheme.*?</a:majorFont>", tx, re.S)
print("=== ШРИФТЫ ===")
print(m2.group(0)[:400] if m2 else "не найдено")