# -*- coding: utf-8 -*-
"""Подстраховка по длине винтов: слайды 6, 14 и комплект (слайд 4)."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
FONT = "Arial"
ORANGE = RGBColor(0xFE, 0xA6, 0x21)
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x5F, 0x5F, 0x5F)
CARD = RGBColor(0xF2, 0xF3, 0xF5)

prs = Presentation(P)
NOTE = ("Не дотягивается М3х10 — ", "возьми М3х20.")
NOTES_TXT = ("Если винт М3х10 не дотягивается и гайка болтается - ставь М3х20 "
             "(в комплекте их 6 шт, на эти два шага нужно 4). "
             "По заявлению преподавателя нужны винты 3х16, но в комплекте их нет - "
             "уточнить у поставщика.")


def para(tf, first=False, runs=None, size=12, color=DARK, bold=False,
         align=PP_ALIGN.LEFT, after=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(after)
    for t, o in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return p


# ---------- 1. Плашки на слайдах 6 и 14 ----------

spots = {6: (Inches(5.95), Inches(4.72), Inches(3.9), Inches(0.72)),
         14: (Inches(0.4), Inches(4.72), Inches(5.4), Inches(0.62))}
for no, (x, y, w, h) in spots.items():
    s = prs.slides[no - 1]
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.fill.background()
    card.shadow.inherit = False
    card.adjustments[0] = 0.12
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.06)
    para(tf, first=True, runs=[
        (NOTE[0], {"bold": True, "color": ORANGE}),
        (NOTE[1], {}),
    ], size=12)
    # заметки докладчика
    ntf = s.notes_slide.notes_text_frame
    ntf.text = ntf.text + "\n" + NOTES_TXT
    print(f"слайд {no}: плашка добавлена")

# ---------- 2. Строка на слайд 4 «Комплект на набор» ----------

s4 = prs.slides[3]
target = None
for sh in s4.shapes:
    if sh.has_text_frame and "разложи детали" in sh.text_frame.text:
        target = sh
        break
tf = target.text_frame
para(tf, runs=[("Если М3х10 не дотягивается (гайка болтается) — ставь М3х20: слайды 6 и 14.", {})],
     size=12, color=GRAY)
print("слайд 4: строка добавлена")
s4.notes_slide.notes_text_frame.text = \
    s4.notes_slide.notes_text_frame.text + "\n" + NOTES_TXT

prs.save(P)
print("\nсохранено")