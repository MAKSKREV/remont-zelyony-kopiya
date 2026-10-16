# -*- coding: utf-8 -*-
"""Ищу все списки компонентов/винтов и смотрю слайды 4, 30, 31."""
import sys, os, re
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
P = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки 4).pptx")
prs = Presentation(P)

print("=== слайды 4, 29, 30, 31 полностью ===")
for no in (4, 29, 30, 31):
    s = prs.slides[no - 1]
    print(f"\n--- слайд {no} layout={s.slide_layout.name} ---")
    for sh in s.shapes:
        kind = {13: "ФОТО", 6: "ГРУППА", 17: "ТЕКСТ", 14: "PH", 1: "ФИГУРА"}.get(sh.shape_type, str(sh.shape_type))
        geo = ""
        if sh.left is not None:
            geo = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
        print(f"  {kind} {geo}")
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.runs or p.text.strip():
                    sz = p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else "насл"
                    print(f"      [{sz}] {p.text[:100]}")

print("\n=== все упоминания винти/винт/гайк/комплект/М3 (видимый текст) ===")
for i, s in enumerate(prs.slides, 1):
    hits = []
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n", " | ")
            if re.search(r"винт|гайк|комплект|М3х", t, re.I):
                hits.append(t[:160])
    if hits:
        print(f"[{i}] " + " ‖ ".join(hits))
