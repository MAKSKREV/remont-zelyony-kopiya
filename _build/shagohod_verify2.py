# -*- coding: utf-8 -*-
import sys, zipfile, os
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
print("слайдов:", len(prs.slides._sldIdLst))
for i, s in enumerate(prs.slides, 1):
    title = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
    if not title:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                title = sh.text_frame.text.strip().split("\n")[0]
                break
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    mark = " <-- новый/правленый" if i in (5, 7, 13, 24, 26, 39, 40, 41, 42) else ""
    print(f"[{i:2}] фото={pics} :: {title[:60]}{mark}")

print("\n=== слайд 24 ===")
for sh in prs.slides[23].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  ", sh.text_frame.text.strip().replace("\n", " / ")[:120])
print("\n=== слайд 26 (новый) ===")
for sh in prs.slides[25].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  ", sh.text_frame.text.strip().replace("\n", " / ")[:200])
print("\n=== слайд 42 (гифка) ===")
for sh in prs.slides[41].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  текст:", sh.text_frame.text.strip())
    if sh.shape_type == 13:
        print("  файл:", sh.image.filename, "|", round(len(sh.image.blob)/1e6, 2), "МБ")

z = zipfile.ZipFile(P)
print("\nразмер файла MB:", round(os.path.getsize(P)/1e6, 2), "| zip:", "ок" if z.testzip() is None else "ОШИБКА")
prs2 = Presentation(P)
print("открывается повторно, слайдов:", len(prs2.slides._sldIdLst))