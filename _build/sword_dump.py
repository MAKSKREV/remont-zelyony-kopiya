# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\5. Инструкция для сборки Светового меча.pptx"
prs = Presentation(P)
print("slides:", len(prs.slides._sldIdLst))
for i, slide in enumerate(prs.slides, 1):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip().replace("\n", " | ")[:90])
    print(f"[{i}] {slide.slide_layout.name} :: {' // '.join(out[:4])}")