# -*- coding: utf-8 -*-
"""Финальная проверка: связи, дубли id, слайд 30, сравнение наложений с исходником."""
import sys, os, re, zipfile
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
OUT = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки 4).pptx")
SRC = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки) 3 (2).pptx")

# ---- 1. все r:embed / r:id слайдов разрешаются в rels?
z = zipfile.ZipFile(OUT)
problems = []
for n in z.namelist():
    if re.match(r"ppt/slides/slide\d+\.xml$", n):
        rels_name = n.replace("slides/", "slides/_rels/") + ".rels"
        xml = z.read(n).decode("utf-8", "ignore")
        rels = z.read(rels_name).decode("utf-8", "ignore")
        used = set(re.findall(r'r:(?:embed|id|link)="([^"]+)"', xml))
        defined = set(re.findall(r'Id="([^"]+)"', rels))
        missing = used - defined
        if missing:
            problems.append((n, missing))
print("битые связи:", problems if problems else "нет")

# ---- 2. дубли id фигур внутри слайдов
prs = Presentation(OUT)
dupes = []
for i, s in enumerate(prs.slides, 1):
    ids = [el.get("id") for el in s.shapes._spTree.iter() if el.tag.endswith("}cNvPr")]
    seen, double = set(), set()
    for x in ids:
        (double if x in seen else seen).add(x)
    if double:
        dupes.append((i, double))
print("дубли id фигур:", dupes if dupes else "нет")

# ---- 3. слайд 30 и 31 целиком
for no in (30, 31):
    s = prs.slides[no - 1]
    print(f"\n===== слайд {no} ===== layout={s.slide_layout.name}")
    for sh in s.shapes:
        kind = {13: "ФОТО", 6: "ГРУППА", 17: "ТЕКСТ", 14: "PH", 1: "ФИГУРА"}.get(sh.shape_type, str(sh.shape_type))
        geo = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
        print(f"  {kind} {geo}")
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.runs:
                    sz = p.runs[0].font.size.pt if p.runs[0].font.size else "насл"
                    print(f"      [{sz}] {p.text[:88]}")
    try:
        print("  заметки:", repr(s.notes_slide.notes_text_frame.text[:250]))
    except Exception:
        pass

# ---- 4. наложения в исходнике (были ли и там)
def overlaps(path, nums):
    p = Presentation(path)
    res = {}
    for no in nums:
        s = p.slides[no - 1]
        bs = []
        for sh in s.shapes:
            if sh.left is None:
                continue
            l, t = Emu(sh.left).inches, Emu(sh.top).inches
            bs.append((sh.name[:16], l, t, l + Emu(sh.width).inches, t + Emu(sh.height).inches))
        hits = []
        for i in range(len(bs)):
            for j in range(i + 1, len(bs)):
                a, b = bs[i], bs[j]
                if min(a[3], b[3]) - max(a[1], b[1]) > 0.05 and min(a[4], b[4]) - max(a[2], b[2]) > 0.05:
                    hits.append((a[0], b[0]))
        res[no] = hits
    return res

print("\nисходник, наложения на 45 и 47 (эталон для сравнения):", overlaps(SRC, (45, 47)))

# ---- 5. медиа не потерялись
zs, zo = zipfile.ZipFile(SRC), zipfile.ZipFile(OUT)
ms = sorted(n for n in zs.namelist() if n.startswith("ppt/media/"))
mo = sorted(n for n in zo.namelist() if n.startswith("ppt/media/"))
print("\nмедиа: исходник", len(ms), "| результат", len(mo),
      "| потеряно:", sorted(set(ms) - set(mo)) or "нет",
      "| добавлено:", sorted(set(mo) - set(ms)) or "нет")
