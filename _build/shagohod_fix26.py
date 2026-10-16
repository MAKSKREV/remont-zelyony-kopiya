# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
s = prs.slides[25]  # слайд 26

KID = "Проверь, что все детали на месте, до начала занятия - так соберешь быстрее."
for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        if not p.runs:
            continue
        if "допишите" in "".join(r.text for r in p.runs):
            p.runs[0].text = KID
            for r in p.runs[1:]:
                r.text = ""
            for r in p.runs:
                r.font.size = Pt(13)
                r.font.color.rgb = RGBColor(0x5F, 0x5F, 0x5F)
            r.font.name = "Arial"
            print("исправлено:", p.runs[0].text)

prs.save(P)
prs2 = Presentation(P)
for sh in prs2.slides[25].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  ", sh.text_frame.text.strip().replace("\n", " / ")[:220])