# -*- coding: utf-8 -*-
"""Шагоход: правки текста + слайды диагностики с фото из этой же инструкции."""
import sys, copy
sys.stdout.reconfigure(encoding="utf-8")
from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

PATH = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"

FONT = "Arial"
ORANGE = RGBColor(0xFE, 0xA6, 0x21)
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x5F, 0x5F, 0x5F)
CARD = RGBColor(0xF2, 0xF3, 0xF5)

prs = Presentation(PATH)
layout = prs.slide_layouts[1]  # CUSTOM_2
BADGE_SRC = prs.slides[36]      # слайд 37 — эталонный «важный момент»


# ---------- утилиты ----------

def set_bullet(p, color="FEA621", indent=0.32):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont", "a:buSzPct"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = pPr.makeelement(qn("a:srgbClr"), {"val": color})
    buClr.append(srgb)
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "80000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": FONT}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u25cf"}))


def para(tf, first=False, text=None, runs=None, size=18, bold=False, color=DARK,
         align=PP_ALIGN.LEFT, after=10, line=None, bullet=False):
    if runs is None:
        runs = [(text or "", {})]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    if bullet:
        set_bullet(p)
    for t, o in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", FONT)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", False)
        f.color.rgb = o.get("color", color)
    return p


def card(slide, x, y, w, h, fill=CARD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.adjustments[0] = 0.06
    sh.text_frame.text = ""
    return sh


def new_slide(title, badge=True):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    s.shapes.title.text_frame.text = title
    if badge:
        for sh in BADGE_SRC.shapes:
            if sh.shape_type == 6:  # группа-маркер, без картинок внутри
                s.shapes._spTree.append(copy.deepcopy(sh._element))
    return s


def pic_blob(slide_no):
    for sh in prs.slides[slide_no - 1].shapes:
        if sh.shape_type == 13:
            return sh.image.blob, sh.width, sh.height
    raise RuntimeError(f"на слайде {slide_no} нет фото")


def add_pic_fit(slide, slide_no, x, y, box_w, box_h, caption):
    blob, w, h = pic_blob(slide_no)
    scale = min(box_w / w, box_h / h)
    pw, ph = int(w * scale), int(h * scale)
    pic = slide.shapes.add_picture(BytesIO(blob), x + int((box_w - pw) / 2),
                                   y + int((box_h - ph) / 2), pw, ph)
    tb = slide.shapes.add_textbox(x, y + box_h + Inches(0.06), box_w, Inches(0.3))
    tb.text_frame.word_wrap = True
    para(tb.text_frame, first=True, text=caption, size=13, color=GRAY,
         align=PP_ALIGN.CENTER, after=0)
    return pic


def set_notes(slide, goal, text, note):
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    para(tf, first=True, text="Цель слайда:", size=12, bold=True, after=2)
    para(tf, text=goal, size=12, after=6)
    para(tf, text="Текст:", size=12, bold=True, after=2)
    para(tf, text=text, size=12, after=6)
    para(tf, text="Примечания:", size=12, bold=True, after=2)
    para(tf, text=note, size=12, after=0)


# ---------- 1. Винты М3х10 -> М3х16 (слайды 5 и 13) ----------

screw_fixed = 0
for no in (5, 13):
    for sh in prs.slides[no - 1].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "М3х10" in r.text:
                    r.text = r.text.replace("М3х10", "М3х16")
                    screw_fixed += 1
print("заменено упоминаний винтов:", screw_fixed)

# ---------- 2. Слайд «проверяем по порядку» ----------

s = new_slide("Шагоход не едет - проверяем по порядку")
c = card(s, Inches(0.35), Inches(0.72), Inches(9.3), Inches(3.15))
tf = c.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.28)
tf.margin_right = Inches(0.22)
tf.margin_top = Inches(0.18)
tf.vertical_anchor = MSO_ANCHOR.TOP
steps = [
    "Батарейки заряжены и вставлены до конца, плюс к плюсу.",
    "Кнопка нажимается, контакты чистые, клей их не залил.",
    "Оба провода подключены к клеммам двигателя, обрыва нет.",
    "Все гайки затянуты и зафиксированы клеем, стенки целые.",
    "Шайбы на валу на месте, лапы крутятся от руки без заеданий.",
]
for i, t in enumerate(steps):
    para(tf, first=(i == 0), runs=[(f"{i + 1}. ", {"color": ORANGE, "bold": True}), (t, {})],
         size=17, after=9, line=1.1)

hint = s.shapes.add_textbox(Inches(0.35), Inches(4.12), Inches(9.3), Inches(0.95))
hint.text_frame.word_wrap = True
para(hint.text_frame, first=True, size=16, bold=True, color=ORANGE, after=4,
     text="Все пункты в порядке? Ставим на ровную поверхность и проверяем ход.")
para(hint.text_frame, size=14, color=GRAY,
     text="Не едет - возвращаемся к списку и проверяем по одному пункту. Несколько проблем "
          "встречаются вместе, поэтому идём по порядку сверху вниз.")
set_notes(s, "Показать порядок поиска неисправности, чтобы не гадать.",
          "Пять проверок: питание, кнопка, провода, крепления, шайбы.",
          "Попросите ребёнка называть каждый пункт вслух и отмечать выполненный. "
          "Чаще всего шагоход не едет из-за батареек, кнопки или ослабшей гайки.")

# ---------- 3. Слайд «частые проблемы» ----------

s = new_slide("Частые проблемы и что делать")
problems = [
    ("Не едет совсем", "Батарейки, кнопка или контакт у двигателя. Проверь пункты 1-3."),
    ("Один мотор молчит", "Обрыв или плохой контакт в клемме. Перепаяй и заизолируй клеем."),
    ("Едет рывками", "Окислились клеммы или батарейка вставлена не до конца."),
    ("Гудит, но еле едет", "Заклинило вал или перекос шайб. Поправь шайбы и зазор."),
    ("Едет криво, цепляет корпус", "Не затянуты гайки или разъехались стенки. Подтяни и подклей."),
    ("Батарейки садятся быстро", "Где-то замыкает кнопка или провод. Отключи провод двигателя."),
]
bw, bh = Inches(4.62), Inches(1.32)
pos = [(Inches(0.34), Inches(0.72)), (Inches(5.04), Inches(0.72)),
       (Inches(0.34), Inches(2.14)), (Inches(5.04), Inches(2.14)),
       (Inches(0.34), Inches(3.56)), (Inches(5.04), Inches(3.56))]
for (title_, fix), (x, y) in zip(problems, pos):
    c = card(s, x, y, bw, bh)
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para(tf, first=True, text=title_, size=15, bold=True, color=ORANGE, after=3)
    para(tf, text=fix, size=13, color=DARK, after=0, line=1.05)
set_notes(s, "Научить узнавать типовые неисправности по симптому.",
          "Шесть самых частых жалоб и что делать в каждом случае.",
          "Таблицу можно распечатать и держать на рабочем столе при сборке.")

# ---------- 4. Слайд с фото из инструкции ----------

s = new_slide("Перед запуском: питание, шайбы и крепления")
cells = [(22, "Батарейный отсек закреплен"),
         (24, "Электроника установлена, провода заизолированы"),
         (28, "Деревянные шайбы на валу двигателя на месте"),
         (34, "Гайки подклеены - не раскрутятся на ходу")]
bw2, bh2 = Inches(4.6), Inches(1.62)
pos = [(Inches(0.4), Inches(0.66)), (Inches(5.15), Inches(0.66)),
       (Inches(0.4), Inches(2.9)), (Inches(5.15), Inches(2.9))]
for (src, cap), (x, y) in zip(cells, pos):
    add_pic_fit(s, src, x, y, bw2, bh2, cap)
set_notes(s, "Показать готовые узлы, которые проверяем перед первым запуском.",
          "Четыре фото из инструкции с короткими подписями.",
          "Фото взяты из слайдов 22, 24, 28 и 34 этой же инструкции. "
          "Можно использовать как памятку - распечатать отдельно.")

# ---------- 5. Перестановка: новые слайды после 37 ----------

sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_ids, rest = ids[-3:], ids[:-3]
order = rest[:37] + new_ids + rest[37:]
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in order:
    sldIdLst.append(el)

prs.save(PATH)
print("сохранено:", PATH)
print("слайдов всего:", len(prs.slides._sldIdLst))
for i, sl in enumerate(prs.slides, 1):
    if i in (5, 6, 7, 8, 13, 36, 37, 38, 39, 40, 41):
        t = sl.shapes.title.text_frame.text if sl.shapes.title is not None else "(нет заголовка)"
        print(f"  [{i}] {t}")