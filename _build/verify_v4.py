# -*- coding: utf-8 -*-
"""Проверка «правки 4»: структура, тексты, шрифты, комментарии, наложения."""
import sys, os, re, zipfile
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
OUT = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки 4).pptx")
SRC = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки) 3 (2).pptx")

z = zipfile.ZipFile(OUT)
bad = z.testzip()
print("zip ok:", bad is None, "| размер:", os.path.getsize(OUT) // 1024, "КБ")

names = z.namelist()
comments = [n for n in names if re.match(r"ppt/comments/comment\d+\.xml", n)]
print("файлов комментариев:", len(comments), "| commentAuthors:", "ppt/commentAuthors.xml" in names)
authors = z.read("ppt/commentAuthors.xml").decode("utf-8", "ignore")
print("автор:", re.search(r'name="([^"]+)"', authors).group(1))
from xml.etree import ElementTree as ET
for n in sorted(comments):
    root = ET.fromstring(z.read(n))
    for cm in root:
        if cm.tag.endswith("cm"):
            t = cm.find("{http://schemas.openxmlformats.org/presentationml/2006/main}text")
            print("   коммент:", "".join(t.itertext()) if t is not None else "?")

prs = Presentation(OUT)
print("\nслайдов:", len(prs.slides))

print("\n=== структура ===")
for i, s in enumerate(prs.slides, 1):
    title = ""
    for sh in s.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 0 and sh.has_text_frame:
            title = sh.text_frame.text.strip().replace("\n", " ")
            break
    if not title:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                title = sh.text_frame.text.strip().replace("\n", " ")[:60]
                break
    print(f"{i:3}. {title}")

print("\n=== текстовые проверки ===")
vis = ""
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            vis += sh.text_frame.text + "\n"
checks = [
    ("М3х16 есть", "М3х16" in vis),
    ("М3х10 нет", "М3х10" not in vis),
    ("М3х20 остался", "М3х20" in vis),
    ("«Эти провода»", "Эти провода длиной" in vis),
    ("заглушек «Добавить гифку» нет", "Добавить гифку" not in vis),
]
for k, v in checks:
    print(("OK  " if v else "FAIL"), k)

refs = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for line in sh.text_frame.text.split("\n"):
                if re.search(r"слайд\w*\s+\d", line, re.I):
                    refs.append((i, line.strip()))
print("ссылки на номера слайдов в видимом тексте:", refs if refs else "нет")

print("\n=== шрифты < 14pt (видимые) ===")
small = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size.pt < 14 and r.text.strip():
                        small.append((i, r.font.size.pt, r.text.strip()[:50]))
print(small if small else "нет")

print("\n=== целевые слайды после перестановки ===")
for no in (5, 13, 16, 17, 18, 19, 28, 29, 49):
    s = prs.slides[no - 1]
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n", " / ")[:70])
    print(f"[{no}] " + " || ".join(texts[:5]))

print("\n=== наложения рамок (ключевые слайды) ===")
def boxes(s):
    out = []
    for sh in s.shapes:
        if sh.left is None:
            continue
        l, t = Emu(sh.left).inches, Emu(sh.top).inches
        r, b = l + Emu(sh.width).inches, t + Emu(sh.height).inches
        out.append((sh.shape_id, sh.name[:18], l, t, r, b))
    return out

def overlap(a, b):
    ix = min(a[4], b[4]) - max(a[2], b[2])
    iy = min(a[5], b[5]) - max(a[3], b[3])
    return ix > 0.05 and iy > 0.05

for no in (16, 17, 18, 28, 47, 49, 50):
    s = prs.slides[no - 1]
    bs = boxes(s)
    hits = []
    for i in range(len(bs)):
        for j in range(i + 1, len(bs)):
            if overlap(bs[i], bs[j]):
                hits.append((bs[i][1], bs[j][1]))
    # выход за границы слайда
    oob = [(b[1], round(b[2], 2), round(b[3], 2), round(b[4], 2), round(b[5], 2))
           for b in bs if b[2] < -0.05 or b[3] < -0.05 or b[4] > 10.05 or b[5] > 5.68]
    print(f"[{no}] наложения={hits or 'нет'} | вне поля={oob or 'нет'}")
