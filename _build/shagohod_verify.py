# -*- coding: utf-8 -*-
import sys, zipfile, os
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

print("=== винты на слайдах 5 и 13 ===")
for no in (5, 13):
    for sh in prs.slides[no-1].shapes:
        if sh.has_text_frame and "Винты" in sh.text_frame.text:
            print(f"  слайд {no}: {sh.text_frame.text.strip().replace(chr(10), ' | ')}")

print("\n=== новые слайды 38-40 ===")
for no in (38, 39, 40):
    s = prs.slides[no-1]
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    groups = sum(1 for sh in s.shapes if sh.shape_type == 6)
    print(f"\n-- слайд {no}: фото={pics}, маркер={groups}, макет={s.slide_layout.name}")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("   ", sh.text_frame.text.strip().replace("\n", " / ")[:200])
    nt = s.notes_slide.notes_text_frame.text.strip().replace("\n", " ")
    print("    ЗАМЕТКИ:", nt[:150])

print("\n=== файл ===")
print("  размер MB:", round(os.path.getsize(P)/1e6, 2))
z = zipfile.ZipFile(P)
print("  изображений в пакете:", len([n for n in z.namelist() if n.startswith("ppt/media/")]))
print("  слайдов в пакете:", len([n for n in z.namelist() if n.startswith("ppt/slides/slide")]))
bad = z.testzip()
print("  целостность zip:", "ок" if bad is None else f"ОШИБКА {bad}")
prs2 = Presentation(P)
print("  открывается повторно, слайдов:", len(prs2.slides._sldIdLst))
print("  слайд 41:", prs2.slides[40].shapes[0].text_frame.text.strip() if prs2.slides[40].shapes[0].has_text_frame else "?")