# -*- coding: utf-8 -*-
import sys, os, zipfile
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
print("слайдов:", len(prs.slides._sldIdLst), "| MB:", round(os.path.getsize(P)/1e6, 2))
z = zipfile.ZipFile(P)
print("zip:", "ок" if z.testzip() is None else "ОШИБКА")

print("\n=== слайды 4, 6, 14 ===")
for no in (4, 6, 14):
    s = prs.slides[no-1]
    print(f"\n-- слайд {no}: {s.shapes.title.text_frame.text if s.shapes.title is not None else ''}")
    for sh in s.shapes:
        L, T = Emu(sh.left).inches, Emu(sh.top).inches
        R, B = L + Emu(sh.width).inches, T + Emu(sh.height).inches
        if sh.has_text_frame and sh.text_frame.text.strip():
            print(f"   текст ({L:.2f},{T:.2f}-{R:.2f},{B:.2f}): "
                  f"{sh.text_frame.text.strip().replace(chr(10), ' / ')[:150]}")

print("\n=== проверка пересечений плашек с фото ===")
for no in (6, 14):
    s = prs.slides[no-1]
    cards = [sh for sh in s.shapes if sh.shape_type == 1 and sh.has_text_frame
             and "дотягивается" in sh.text_frame.text]
    for c in cards:
        L, T = Emu(c.left).inches, Emu(c.top).inches
        R, B = L + Emu(c.width).inches, T + Emu(c.height).inches
        hits = []
        for sh in s.shapes:
            if sh.shape_type != 13:
                continue
            l2, t2 = Emu(sh.left).inches, Emu(sh.top).inches
            r2, b2 = l2 + Emu(sh.width).inches, t2 + Emu(sh.height).inches
            if L < r2 and R > l2 and T < b2 and B > t2:
                hits.append((round(l2, 2), round(t2, 2), round(r2, 2), round(b2, 2)))
        print(f"  слайд {no}: плашка ({L:.2f},{T:.2f}-{R:.2f},{B:.2f}) -> фото внутри: {len(hits)} {hits if hits else ''}")

prs2 = Presentation(P)
print("\nоткрывается повторно, слайдов:", len(prs2.slides._sldIdLst))