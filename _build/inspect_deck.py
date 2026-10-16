# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
prs = Presentation(r"C:\Users\Maks\Desktop\ремонтзелёный  копия\TRIZ-master-klass.pptx")
print("slides:", len(prs.slides._sldIdLst))
for i, slide in enumerate(prs.slides, 1):
    t = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t.append(sh.text_frame.text.strip().replace("\n", " / ")[:70])
        if len(t) >= 2:
            break
    print(i, slide.slide_layout.name, "|", " || ".join(t))