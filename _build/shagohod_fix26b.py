# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
s = prs.slides[25]

ORANGE = RGBColor(0xFE, 0xA6, 0x21)
GRAY = RGBColor(0x5F, 0x5F, 0x5F)

NEW = {
    "Перед занятием разложи все детали по группам - так соберешь быстрее.":
        ("Разложи детали по группам перед занятием - так соберешь быстрее.", ORANGE, 15, True),
    "Проверь, что все детали на месте, до начала занятия - так соберешь быстрее.":
        ("Шагоход не поедет? Смотри разбор неисправностей на слайдах 39-41.", GRAY, 13, False),
}

for sh in s.shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        cur = "".join(r.text for r in p.runs)
        if cur in NEW:
            text, color, size, bold = NEW[cur]
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ""
            r = p.runs[0]
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Arial"
            print("->", r.text)

prs.save(P)
prs2 = Presentation(P)
print()
for sh in prs2.slides[25].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  ", sh.text_frame.text.strip().replace("\n", " / ")[:230])