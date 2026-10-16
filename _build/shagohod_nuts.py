# -*- coding: utf-8 -*-
"""Гайки на слайдах 28, 34, 37 + обновление списка деталей на занятие 2."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
FONT, ORANGE, GRAY, DARK = "Arial", RGBColor(0xFE, 0xA6, 0x21), RGBColor(0x5F, 0x5F, 0x5F), RGBColor(0x21, 0x21, 0x21)
prs = Presentation(P)


def set_bullet(p, color="FEA621", indent=0.32):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    buClr = pPr.makeelement(qn("a:buClr"), {})
    buClr.append(pPr.makeelement(qn("a:srgbClr"), {"val": color}))
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "80000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": FONT}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u25cf"}))


def para(tf, first=False, text=None, runs=None, size=17, bold=False, color=DARK,
         align=PP_ALIGN.LEFT, after=10, line=None, bullet=False):
    if runs is None:
        runs = [(text or "", {})]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    if bullet:
        set_bullet(p)
    for t, o in runs:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", FONT)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return p


# ---------- 1. Подписи с гайками на слайдах 28, 34, 37 ----------

targets = {
    28: (Inches(7.48), Inches(1.05), 2, "2шт."),
    34: (Inches(0.18), Inches(1.05), 2, "3шт."),
    37: (Inches(0.18), Inches(1.05), 2, "3шт."),
}
for no, (x, y, w, qty) in targets.items():
    s = prs.slides[no - 1]
    tb = s.shapes.add_textbox(x, y, Inches(w), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    para(tf, first=True, text=f"Гайки М3 - {qty}", size=16, bold=True, color=DARK, after=0)
    print(f"слайд {no}: Гайки М3 - {qty}")

# ---------- 2. Список деталей на слайде 26 ----------

s26 = prs.slides[25]
items = [
    "Винты М3х20 - 2 шт. и гайки М3 - 2 шт. (детали, слайд 31)",
    "Гайки М3 - 8 шт.: стенка - 2, корпус - 3, вторая деталь - 3",
    "Деревянные шайбы - на 3 места: вал двигателя, крепление, вторая деталь",
    "Электроника с занятия 1: батарейки, кнопка, двигатель, провода",
    "Клей и инструмент - как на занятии 1",
]
target = None
for sh in s26.shapes:
    if sh.has_text_frame and "Винты М3х20" in sh.text_frame.text:
        target = sh
        break
tf = target.text_frame
tf.clear()
for i, t in enumerate(items):
    para(tf, first=(i == 0), text=t, size=17, after=9, line=1.1, bullet=True)

# ---------- 3. Заметки к слайду 26 ----------

ntf = s26.notes_slide.notes_text_frame
ntf.text = ""
para(ntf, first=True, text="Цель слайда:", size=12, bold=True, after=2)
para(ntf, text="Дать ребенку заранее подготовить все детали на занятие 2.", size=12, after=6)
para(ntf, text="Текст:", size=12, bold=True, after=2)
para(ntf, text="Список деталей перед началом занятия: 10 гаек М3, винты М3х20 - 2 шт., "
               "деревянные шайбы, электроника, клей.", size=12, after=6)
para(ntf, text="Примечания:", size=12, bold=True, after=2)
para(ntf, text="Количество гаек уточнено: 2 - слайд 28, 3 - слайд 34, 3 - слайд 37. "
               "Длина винтов на этих слайдах в инструкции не указана - допишите, если нужно.",
    size=12, after=0)

prs.save(P)
print("\nсохранено")
for i in (26, 28, 34, 37):
    s = prs.slides[i-1]
    print(f"\n-- слайд {i}: {s.shapes.title.text_frame.text if s.shapes.title is not None else ''}")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("   ", sh.text_frame.text.strip().replace("\n", " / ")[:190])