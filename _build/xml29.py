# -*- coding: utf-8 -*-
"""XML плейсхолдеров и фото слайда 29 - эталон для слайда 28."""
import sys, os
sys.stdout.reconfigure(encoding="utf-8")
from lxml import etree
from pptx import Presentation

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
P = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки) 3 (2).pptx")
prs = Presentation(P)
s = prs.slides[28]
for sh in s.shapes:
    if sh.shape_id in (430, 431, 432):
        print("=" * 30, sh.shape_id, sh.name)
        xml = etree.tostring(sh._element, pretty_print=True).decode()
        print(xml[:5000])
