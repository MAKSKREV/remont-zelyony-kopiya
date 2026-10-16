# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

for no in (28, 31, 34, 37):
    s = prs.slides[no-1]
    print(f"\n=== слайд {no}: {s.shapes.title.text_frame.text if s.shapes.title is not None else ''}")
    for sh in s.shapes:
        kind = {13: "ФОТО", 6: "ГРУППА", 17: "ТЕКСТ", 1: "ФИГУРА", 14: "ПЛЕЙСХОЛДЕР"}.get(sh.shape_type, str(sh.shape_type))
        L, T = Emu(sh.left).inches, Emu(sh.top).inches
        R, B = L + Emu(sh.width).inches, T + Emu(sh.height).inches
        txt = sh.text_frame.text.strip().replace("\n", " / ")[:45] if sh.has_text_frame else ""
        sz = ""
        if sh.has_text_frame and sh.text_frame.paragraphs and sh.text_frame.paragraphs[0].runs:
            f = sh.text_frame.paragraphs[0].runs[0].font
            sz = f" {f.size.pt if f.size else '?'}pt bold={f.bold}"
        print(f"  {kind:11} x:{L:5.2f}-{R:5.2f}  y:{T:5.2f}-{B:5.2f}{sz}  {txt}")