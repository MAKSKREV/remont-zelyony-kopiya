# -*- coding: utf-8 -*-
import sys, zipfile, os
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)
print("слайдов:", len(prs.slides._sldIdLst), "| размер MB:", round(os.path.getsize(P)/1e6, 2))
z = zipfile.ZipFile(P)
print("zip:", "ок" if z.testzip() is None else "ОШИБКА",
      "| изображений:", len([n for n in z.namelist() if n.startswith("ppt/media/")]))

print("\n=== проверка, что подписи не налезают на фото ===")
for no, box in ((28, (7.48, 1.05, 2.0, 0.55)), (34, (0.18, 1.05, 2.0, 0.55)), (37, (0.18, 1.05, 2.0, 0.55))):
    s = prs.slides[no-1]
    bx = (box[0], box[1], box[0]+box[2], box[1]+box[3])
    bad = []
    for sh in s.shapes:
        if sh.shape_type != 13:
            continue
        L, T = Emu(sh.left).inches, Emu(sh.top).inches
        R, B = L + Emu(sh.width).inches, T + Emu(sh.height).inches
        if bx[0] < R and bx[2] > L and bx[1] < B and bx[3] > T:
            bad.append((round(L,2), round(T,2), round(R,2), round(B,2)))
    print(f"  слайд {no}: пересечений с фото: {len(bad)} {bad if bad else ''}")

print("\n=== итоговые слайды занятия 2 ===")
for i in range(25, 39):
    s = prs.slides[i-1]
    t = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
    extra = [sh.text_frame.text.strip().replace("\n", " / ") for sh in s.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()
             and "Гайки" in sh.text_frame.text and sh.shape_type == 17]
    print(f"[{i:2}] {t[:55]:57} {extra if extra else ''}")
prs2 = Presentation(P)
print("\nоткрывается повторно, слайдов:", len(prs2.slides._sldIdLst))