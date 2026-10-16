# -*- coding: utf-8 -*-
"""Сводная таблица: где и сколько гаек и винтов указано."""
import sys, re
from collections import defaultdict
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

rows = []
for i, s in enumerate(prs.slides, 1):
    t = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in p.runs).strip()
            m = re.match(r"^(Винты|Гайки)\s+(М3\S*)\s*-\s*(\d+)\s*шт", txt, re.I)
            if m:
                rows.append((i, t, m.group(1), m.group(2), int(m.group(3))))

print("ВСЕ УКАЗАНИЯ В ФАЙЛЕ (44 слайда)")
print("=" * 92)
cur = None
for i, t, kind, size, qty in rows:
    zone = ("КОМПЛЕКТ" if i == 4 else "ЗАНЯТИЕ 1" if i <= 26 else "ЗАНЯТИЕ 2" if i <= 39 else "ДИАГНОСТИКА")
    if zone != cur:
        print(f"\n--- {zone} ---")
        cur = zone
    print(f"  слайд {i:2}  |  {kind:6} {size:6} {qty:2} шт  |  {t[:48]}")

print("\n" + "=" * 92)
print("ИТОГИ ПО СБОРКЕ (без слайда 4 «Комплект» и слайда 27 «список на занятие»)")
tot = defaultdict(int)
for i, t, kind, size, qty in rows:
    if i in (4, 27):
        continue
    tot[(kind, size)] += qty
for k in sorted(tot):
    print(f"  {k[0]:6} {k[1]:7} {tot[k]:2} шт")

z1 = [r for r in rows if r[0] <= 26 and r[0] != 4]
z2 = [r for r in rows if 27 < r[0] <= 39]
print("\n  Занятие 1:", ", ".join(f"{k} {v} шт" for (kk, sz), v in
      sorted({(k, sz): 0 for _, _, k, sz, _ in z1}.items()) if True) or "")
t1 = defaultdict(int)
for i, t, kind, size, qty in z1:
    t1[(kind, size)] += qty
print("           ", " | ".join(f"{k} {sz}: {v}" for (k, sz), v in sorted(t1.items())))
t2 = defaultdict(int)
for i, t, kind, size, qty in z2:
    t2[(kind, size)] += qty
print("  Занятие 2:", " | ".join(f"{k} {sz}: {v}" for (k, sz), v in sorted(t2.items())))

KIT = {("Винты", "М3х10"): 4, ("Винты", "М3х20"): 6, ("Винты", "М3х30"): 2, ("Гайки", "М3"): 23}
print("\n" + "=" * 92)
print("СВЕРКА С КОМПЛЕКТОМ НА НАБОР")
for k, v in KIT.items():
    need = tot.get(k, 0)
    d = v - need
    flag = "хватает" if d >= 0 else f"НЕ ХВАТАЕТ {-d}"
    print(f"  {k[0]:6} {k[1]:7} нужно {need:2}, в комплекте {v:2}  ({d:+d})  {flag}")