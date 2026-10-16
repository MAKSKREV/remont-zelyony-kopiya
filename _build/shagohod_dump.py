# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход.pptx"
prs = Presentation(P)
print("slides:", len(prs.slides._sldIdLst), "| size:", prs.slide_width, prs.slide_height)
for i, slide in enumerate(prs.slides, 1):
    texts, pics = [], 0
    for sh in slide.shapes:
        if sh.shape_type == 13:
            pics += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n", " | "))
    head = texts[0][:70] if texts else "(без текста)"
    print(f"[{i}] pics={pics} :: {head}")
    for t in texts[1:]:
        print("      -", t[:120])