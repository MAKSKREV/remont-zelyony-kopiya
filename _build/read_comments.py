# -*- coding: utf-8 -*-
"""Комментарии в вернувшейся презентации."""
import sys, os, zipfile, re
from xml.etree import ElementTree as ET
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
P = os.path.join(BASE, "Инструкция по сборке (фото) - шагоход (правки) 3 (2).pptx")
z = zipfile.ZipFile(P)
names = z.namelist()

print("=== части с комментариями ===")
for n in sorted(names):
    if "comment" in n.lower() and not n.endswith(".rels"):
        print("  ", n, z.getinfo(n).file_size)

authors = {}
if "ppt/authors.xml" in names:
    root = ET.fromstring(z.read("ppt/authors.xml"))
    for el in root:
        authors[el.get("id")] = el.get("name") or "?"
    print("\nавторы:", authors)

slide_rels = {}
for n in names:
    m = re.match(r"ppt/slides/_rels/slide(\d+)\.xml\.rels", n)
    if m:
        txt = z.read(n).decode("utf-8", "ignore")
        r = re.search(r'Target="\.\./(comments[^"]*?\.xml)"', txt)
        if r:
            slide_rels[int(m.group(1))] = "ppt/" + r.group(1)
print("\nслайды с комментариями:", sorted(slide_rels))

print("\n=== КОММЕНТАРИИ ===")
for no, rel in sorted(slide_rels.items()):
    tree = ET.fromstring(z.read(rel))
    for cm in tree:
        tag = cm.tag.split("}")[-1]
        if not tag.lower().startswith("comment"):
            continue
        cid = cm.get("id")
        who = authors.get(cid, cm.get("author", "?"))
        date = cm.get("dt", cm.get("date", ""))
        body = " ".join("".join(t.itertext()) for t in cm.iter() if t.tag.endswith("}t")).strip()
        print(f"\n[слайд {no}] {who} {date} id={cid}")
        print("   ", body)

# привязка комментариев к конкретным формам слайда
print("\n=== привязка к формам ===")
for no, rel in sorted(slide_rels.items()):
    txt = z.read(rel).decode("utf-8", "ignore")
    for m in re.finditer(r"<p:cm\w*:commentRangeStart[^>]*?id=\"(\d+)\"[^>]*/>", txt):
        print(f"  слайд {no}: commentRangeStart id={m.group(1)} (позиция в xml: {m.start()})")

prs = Presentation(P)
print("\nслайдов:", len(prs.slides._sldIdLst))
print("\n=== структура ===")
for i, s in enumerate(prs.slides, 1):
    t = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
    if not t:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = "[" + sh.text_frame.text.strip().split("\n")[0][:38] + "]"
                break
    print(f"[{i:2}] {t[:58]}")