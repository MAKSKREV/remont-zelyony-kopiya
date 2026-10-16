# -*- coding: utf-8 -*-
"""Сверяем, с каких слайдов реально взяты фото на слайде 42, по хэшам картинок."""
import sys, hashlib
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

def title(i):
    s = prs.slides[i-1]
    return s.shapes.title.text_frame.text if s.shapes.title is not None else "(без заголовка)"

print("фото на слайде 42:")
for sh in prs.slides[41].shapes:
    if sh.shape_type != 13:
        continue
    h = hashlib.sha1(sh.image.blob).hexdigest()
    src = []
    for i, s in enumerate(prs.slides, 1):
        if i == 42:
            continue
        for other in s.shapes:
            if other.shape_type == 13 and hashlib.sha1(other.image.blob).hexdigest() == h:
                src.append(i)
    print(f"  фото -> источник: слайд {src} «{title(src[0]) if src else '?'}»")

print("\nзаметки слайда 42:")
print(" ", prs.slides[41].notes_slide.notes_text_frame.text.replace("\n", " | "))