# -*- coding: utf-8 -*-
"""Сверка комплекта с тем, что указано на слайдах."""
import sys, re
from collections import defaultdict
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

use = defaultdict(int)
lines = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            m = re.match(r"^(Винты|Гайки)\s+(М3\S*)\s*-\s*(\d+)\s*шт", t, re.I)
            if m:
                kind, size, qty = m.group(1), m.group(2), int(m.group(3))
                use[(kind, size)] += qty
                lines.append((i, kind, size, qty))

print("=== что указано на слайдах ===")
for i, kind, size, qty in lines:
    print(f"  слайд {i:2}: {kind:6} {size:7} {qty} шт")

print("\n=== итог по инструкции ===")
for (kind, size), qty in sorted(use.items()):
    print(f"  {kind:6} {size:7} {qty} шт")

KIT = {("Винты", "М3х20"): 6, ("Винты", "М3х30"): 2, ("Винты", "М3х10"): 4}
print("\n=== комплект на набор ===")
for k, v in KIT.items():
    print(f"  {k[0]:6} {k[1]:7} {v} шт")

print("\n=== сверка (винти) ===")
z1 = 0
for size in ("М3х10", "М3х16", "М3х20", "М3х30"):
    need = use.get(("Винты", size), 0)
    have = KIT.get(("Винты", size), 0)
    if need or have:
        flag = "ок" if need == have else ("НЕ ХВАТАЕТ" if need > have else "остаток")
        print(f"  {size}: нужно {need}, в наборе {have} -> {flag}")
        z1 += max(0, need - have)
print("  всего не хватает:", z1)

print("\n=== гайки ===")
print("  нужно по инструкции:", use.get(("Гайки", "М3"), 0), "шт; в наборе не указаны")