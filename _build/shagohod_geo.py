# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

P = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\Копия Инструкция по сборке (фото) - шагоход (правки).pptx"
prs = Presentation(P)

for no in (17, 24, 25, 26, 27, 33, 36, 41):
    s = prs.slides[no-1]
    print(f"\n=== слайд {no} (макет {s.slide_layout.name}) ===")
    for sh in s.shapes:
        kind = {13: "ФОТО", 6: "ГРУППА", 17: "ТЕКСТ", 1: "ФИГУРА", 14: "ПЛЕЙСХОЛДЕР"}.get(sh.shape_type, str(sh.shape_type))
        txt = sh.text_frame.text.strip().replace("\n", " / ")[:60] if sh.has_text_frame else ""
        print(f"  {kind:11} ({Emu(sh.left).inches:5.2f},{Emu(sh.top).inches:5.2f}) "
              f"{Emu(sh.width).inches:4.2f}x{Emu(sh.height).inches:4.2f}  {txt}")