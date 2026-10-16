# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
OUT = r"C:\Users\Maks\Desktop\ремонтзелёный  копия\ТРИЗ-мастер-класс (в стиле Пример).pptx"
prs = Presentation(OUT)
print("slides:", len(prs.slides._sldIdLst))
for i, slide in enumerate(prs.slides, 1):
    texts = []
    shapes = slide.shapes
    for sh in shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.replace("\n", " / ")
            if t.strip():
                texts.append(t)
    print(f"[{i}] layout={slide.slide_layout.name}")
    for t in texts[:6]:
        print("    -", t[:110])