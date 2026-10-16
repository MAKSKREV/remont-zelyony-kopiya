# -*- coding: utf-8 -*-
"""Write speaker notes from подсказки.pptx comments into the deck files."""
import sys, os
import zipfile
from lxml import etree
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")

BASE = r"C:\Users\Maks\Desktop\ремонтзелёный  копия"
PODSK = os.path.join(BASE, "подсказки.pptx")
TARGETS = [os.path.join(BASE, "TRIZ-master-klass.pptx"),
           os.path.join(BASE, "TRIZ-master-klass для преподавателя.pptx")]

NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def read_comment_texts(path):
    z = zipfile.ZipFile(path)
    names = set(z.namelist())
    slide_to_comments = {}
    for sl in range(1, 100):
        rp = f"ppt/slides/_rels/slide{sl}.xml.rels"
        if rp not in names:
            continue
        rel = etree.fromstring(z.read(rp))
        for r in rel:
            if r.get("Type", "").endswith("/comments"):
                tgt = r.get("Target")
                if tgt.startswith("../"):
                    tgt = "ppt/" + tgt[3:]
                texts = []
                root = etree.fromstring(z.read(tgt))
                for cm in root:
                    node = cm.find(f"{{{NS_P}}}text")
                    if node is not None and node.text:
                        texts.append(node.text)
                slide_to_comments[sl] = texts
    z.close()
    return slide_to_comments


def set_notes(slide, lines):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.text = ""
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line


slide_notes = read_comment_texts(PODSK)
print("slides with comments in подсказки.pptx:", len(slide_notes))

for tgt in TARGETS:
    prs = Presentation(tgt)
    for k, slide in enumerate(prs.slides, 1):
        if k in slide_notes:
            joined = "\n".join(slide_notes[k])
            set_notes(slide, joined.split("\n"))
    prs.save(tgt)
    print("notes rewritten:", os.path.basename(tgt))

# verify
prs = Presentation(TARGETS[0])
zs = 0
for k, slide in enumerate(prs.slides, 1):
    t = slide.notes_slide.notes_text_frame.text
    if t.strip():
        zs += 1
print("slides with non-empty notes:", zs)
print("slide18 notes:")
print(prs.slides[17].notes_slide.notes_text_frame.text)