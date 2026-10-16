# -*- coding: utf-8 -*-
import sys, zipfile
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
for P in (r"C:\Users\Maks\Desktop\ремонтзелёный  копия\TRIZ-master-klass.pptx",
          r"C:\Users\Maks\Desktop\ремонтзелёный  копия\TRIZ-master-klass для преподавателя.pptx"):
    prs = Presentation(P)
    z = zipfile.ZipFile(P)
    nparts = [n for n in z.namelist() if n.startswith("ppt/notesSlides/")]
    mc = [n for n in z.namelist() if n.startswith("ppt/comments/") or "commentAuthor" in n]
    print(P.split("\\")[-1])
    print("  slides:", len(prs.slides._sldIdLst), "| notesSlide parts:", len(nparts),
          "| comment parts:", len(mc))
    ns = prs.slides[8].notes_slide.notes_text_frame
    print("  slide9 notes paragraphs:", len(ns.paragraphs))
    print("  slide9 notes first 2 lines:")
    for p in ns.paragraphs[:2]:
        if p.text.strip():
            print("    -", p.text[:80])
    z.close()